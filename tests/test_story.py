import json
from pathlib import Path

import pytest
from openpyxl import load_workbook

from services import story
from services.qc import QCSettings
from services.story import (
    MAX_LINE_CHARS,
    MAX_STORY_LINES,
    STORY_EMOTIONS,
    StoryLine,
    StorySplitError,
    split_story,
)

SHEET = Path(__file__).resolve().parent.parent / "the_boy_who_followed_the_stream_voiceoversheet.xlsx"


def _reference_lines() -> list[tuple[str, str]]:
    ws = load_workbook(SHEET).active
    rows = list(ws.iter_rows(values_only=True))[1:]
    return [(text, emotion) for text, emotion, *_ in rows]


REFERENCE = _reference_lines()
# The operator pastes plain text; the sheet's cells joined with spaces is that text.
STORY = " ".join(text for text, _ in REFERENCE)

SETTINGS = QCSettings(api_key="test-key", models=["model-a", "model-b"], enabled=True)


def _reply(lines: list[tuple[str, str]]) -> str:
    return json.dumps({"lines": [{"text": t, "emotion": e} for t, e in lines]})


class FakeGemini:
    """Hands out one queued reply per call; records what was asked."""

    def __init__(self, monkeypatch, replies: list[str]):
        self.replies = list(replies)
        self.calls: list[dict] = []
        fake = self

        class FakeResponse:
            def __init__(self, text: str):
                self.text = text

        class FakeModels:
            def generate_content(self, model: str, contents: str, config=None):
                fake.calls.append({
                    "model": model,
                    "contents": contents,
                    "system": getattr(config, "system_instruction", "") or "",
                })
                return FakeResponse(fake.replies.pop(0))

        class FakeClient:
            def __init__(self, api_key: str):
                self.models = FakeModels()

        monkeypatch.setenv("API_RETRY_MAX_ATTEMPTS", "1")
        monkeypatch.setattr("services.story.genai.Client", FakeClient)


def test_happy_split_returns_the_lines_and_emotions(monkeypatch) -> None:
    gemini = FakeGemini(monkeypatch, [_reply(REFERENCE)])

    lines = split_story(STORY, settings=SETTINGS)

    assert lines == [StoryLine(text=t, emotion=e) for t, e in REFERENCE]
    assert len(gemini.calls) == 1
    # the story goes to Gemini, and the allowed emotions are spelled out for it
    assert "Moti wagged his tail and barked." in gemini.calls[0]["contents"]
    for emotion in STORY_EMOTIONS:
        assert emotion in gemini.calls[0]["system"]


def test_quote_style_and_whitespace_do_not_count_as_changed_wording(monkeypatch) -> None:
    # Gemini straightens the curly apostrophe in "Let’s" and the source has
    # stray line breaks; neither is a change to what gets voiced.
    source = "Hi there!   Let’s follow Kiran.\n\n“Moti, look,” said Kiran."
    FakeGemini(monkeypatch, [_reply([
        ("Hi there! Let's follow Kiran.", "Cheerful"),
        ('"Moti, look," said Kiran.', "Narrative"),
    ])])

    lines = split_story(source, settings=SETTINGS)

    assert [line.emotion for line in lines] == ["Cheerful", "Narrative"]


def test_emotion_case_and_typography_are_not_held_against_the_split(monkeypatch) -> None:
    source = "Moti barked… and ran."
    FakeGemini(monkeypatch, [_reply([("Moti barked... and ran.", "narrative")])])

    assert split_story(source, settings=SETTINGS) == [StoryLine("Moti barked... and ran.", "Narrative")]


def test_a_model_failure_falls_back_to_the_next_model(monkeypatch) -> None:
    gemini = FakeGemini(monkeypatch, [_reply(REFERENCE)])
    calls: list[str] = []

    class FlakyModels:
        def generate_content(self, model: str, contents: str, config=None):
            calls.append(model)
            if model == "model-a":
                raise RuntimeError("429 rate limit")
            return type("R", (), {"text": gemini.replies.pop(0)})()

    class FlakyClient:
        def __init__(self, api_key: str):
            self.models = FlakyModels()

    monkeypatch.setattr("services.story.genai.Client", FlakyClient)

    assert len(split_story(STORY, settings=SETTINGS)) == len(REFERENCE)
    assert calls == ["model-a", "model-b"]


def test_changed_wording_is_retried_once_then_rejected(monkeypatch) -> None:
    reworded = list(REFERENCE)
    reworded[6] = ("Moti wagged his tail and barked happily.", "Narrative")
    gemini = FakeGemini(monkeypatch, [_reply(reworded), _reply(reworded)])

    with pytest.raises(StorySplitError, match="wording"):
        split_story(STORY, settings=SETTINGS)

    assert len(gemini.calls) == 2


def test_a_bad_first_reply_is_recovered_by_the_retry(monkeypatch) -> None:
    dropped = REFERENCE[:-1]
    gemini = FakeGemini(monkeypatch, [_reply(dropped), _reply(REFERENCE)])

    lines = split_story(STORY, settings=SETTINGS)

    assert len(lines) == len(REFERENCE)
    assert len(gemini.calls) == 2


def test_unknown_emotion_is_rejected(monkeypatch) -> None:
    bad = list(REFERENCE)
    bad[2] = (bad[2][0], "Melancholic")
    FakeGemini(monkeypatch, [_reply(bad), _reply(bad)])

    with pytest.raises(StorySplitError, match="Melancholic"):
        split_story(STORY, settings=SETTINGS)


def test_a_line_over_the_cap_is_rejected(monkeypatch) -> None:
    # One unbroken run of words: joined it still matches the source, so only the
    # length rule can catch it.
    sentence = "The stream ran on and on past the trees. "
    source = (sentence * (MAX_LINE_CHARS // len(sentence) + 1)).strip()
    FakeGemini(monkeypatch, [_reply([(source, "Narrative")])] * 2)

    with pytest.raises(StorySplitError, match=str(MAX_LINE_CHARS)):
        split_story(source, settings=SETTINGS)


def test_the_cap_counts_the_emotion_tag() -> None:
    text = "a" * (MAX_LINE_CHARS - len("[Narrative] "))
    assert story.line_problem(StoryLine(text=text, emotion="Narrative")) is None
    assert story.line_problem(StoryLine(text=text + "a", emotion="Narrative")) is not None


def test_a_story_split_into_too_many_lines_is_refused(monkeypatch) -> None:
    words = [f"Word{i}." for i in range(MAX_STORY_LINES + 1)]
    source = " ".join(words)
    gemini = FakeGemini(monkeypatch, [_reply([(w, "Narrative") for w in words])] * 2)

    with pytest.raises(StorySplitError, match="too long"):
        split_story(source, settings=SETTINGS)

    # an over-split is a bad split like any other, so it gets the retry too
    assert len(gemini.calls) == 2


def test_a_story_too_big_to_fit_is_refused_without_calling_gemini(monkeypatch) -> None:
    gemini = FakeGemini(monkeypatch, [])
    source = "word " * (MAX_STORY_LINES * MAX_LINE_CHARS // 5 + 1)

    with pytest.raises(StorySplitError, match="too long"):
        split_story(source, settings=SETTINGS)

    assert gemini.calls == []


def test_empty_text_is_refused(monkeypatch) -> None:
    gemini = FakeGemini(monkeypatch, [])

    with pytest.raises(StorySplitError):
        split_story("   \n ", settings=SETTINGS)

    assert gemini.calls == []


def test_reply_that_is_not_json_is_retried_then_rejected(monkeypatch) -> None:
    FakeGemini(monkeypatch, ["Sure! Here are your lines:", "```json\nnope\n```"])

    with pytest.raises(StorySplitError):
        split_story(STORY, settings=SETTINGS)


def test_reply_in_a_code_fence_is_accepted(monkeypatch) -> None:
    FakeGemini(monkeypatch, ["```json\n" + _reply(REFERENCE) + "\n```"])

    assert len(split_story(STORY, settings=SETTINGS)) == len(REFERENCE)


def test_activity_name_from_title() -> None:
    assert story.activity_name("The boy who followed the stream") == "theboywhofollowedthestreamstory"
    assert story.activity_name("  Kiran's Map! (Part 2) ") == "kiransmappart2story"
