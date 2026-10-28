from __future__ import annotations

import asyncio
import re
from io import BytesIO
from contextlib import asynccontextmanager
from dataclasses import dataclass
import logging
from pathlib import Path
import uuid
import tempfile

from dotenv import load_dotenv
from fastapi import FastAPI, File, Form, HTTPException, Request, Response, UploadFile
from fastapi.responses import FileResponse, JSONResponse, RedirectResponse
from fastapi.staticfiles import StaticFiles

from api.background import spawn
from api.logging import get_important_logs as _get_important_logs, install_log_handler
from api.models import (
    ElevenLabsTTSRequest,
    FileRenderRequest,
    FinalizeTextRequest,
    SessionEnvConfigRequest,
    SessionEnvConfigResponse,
    StoryJobRequest,
    StorySplitRequest,
    StoryXlsxRequest,
    TranslateRequest,
)
from api.session_config import SESSION_COOKIE_NAME, SessionConfigStore
from api.utils import ensure_file_extension, parse_target_languages, safe_stem, save_upload_file, to_output_url
from batch.excel import REQUIRED_HEADERS
from batch.models import CreateJobResponse, JobState
from batch.service import english_voice_is_required, run_excel_batch_job
from batch.story_job import run_story_job
from services.elevenlabs import ElevenLabsSettings, ElevenLabsTTSConfig, synthesize_speech_bytes
from batch.store import JobsStore
from services.video_pipeline import (
    NotRecoverable,
    VideoJobsStore,
    run_video_job,
    start_recovery,
    start_recovery_of_failed,
)
from services.video_pipeline.library import page_of
from services.video_pipeline.library_service import Download, LibraryService, NoMedia, sniff_image_type
from services.video_pipeline.pipeline import start_filing_render
from services.video_pipeline.types import VideoJobSpec, VideoJobState
from services.email import EmailSettings
from services.nas import NasConfig, NasService
from services.script_history import ScriptHistoryStore
from services.script_writer import (
    ScriptRequestError,
    ScriptWriter,
    ScriptWriterError,
    ScriptWriterSettings,
    UnknownItemTitleError,
)
from services.video_pipeline.batch_excel import (
    BatchExcelError,
    read_heygen_batch_rows,
    read_heygen_batch_rows_json,
)
from services.video_pipeline.batch_runner import run_video_batch_job
from services.video_pipeline.batch_store import BatchRowState, VideoBatchJobsStore
from services.video_pipeline.control import JobControls
from services.video_pipeline.heygen_client import HeyGenSettings
from services.video_pipeline.heygen_renderer import HeyGenRenderer
from services.video_pipeline.intake import (
    BatchSettings,
    IntakeError,
    SingleSubmission,
    scrub_title,
    spec_from_single,
    specs_from_batch,
)
from services.video_pipeline.renderer import VideoRenderer
from services.video_pipeline.slots import TalkingPhotoSlots
from services.video_pipeline.speech import ElevenLabsSpeech, SpeechSynth
from services.google_translate import GoogleTranslateSettings
from services.sarvam import SarvamSettings
from services.qc import QCSettings
from services.settings import Settings, missing_keys, required_keys
from services.story import (
    MAX_LINE_CHARS,
    STORY_EMOTIONS,
    StoryGeminiError,
    StoryLine,
    StorySplitError,
    activity_name as story_activity_name,
    line_problem,
    split_story,
)
from services.stt import transcribe_audio
from services.tts import text_to_speech
from services.translation import translate_with_fallback
from services.s3 import validate_s3_env
from services.runtime_config import MissingSettingError, parse_env_text, RuntimeConfig


load_dotenv()
logging.basicConfig(
    level=logging.INFO,
    format="%(asctime)s %(levelname)s %(name)s: %(message)s",
)
install_log_handler()


@asynccontextmanager
async def lifespan(app: FastAPI):
    _, s3_error = validate_s3_env()
    app.state.s3_config_error = s3_error
    try:
        NasService(NasConfig.resolve()).ensure_base_folders()
    except Exception as exc:
        logging.warning("NAS base folder init failed (non-fatal): %s", exc)
    # The library's first full listing of the HeyGen account is ~30 requests.
    # Make it now, so the first person to open the library doesn't wait on it.
    try:
        spawn(video_library.catalog.renders(_build_renderer(None)), name="library-warmup")
    except MissingSettingError:
        pass
    yield


app = FastAPI(lifespan=lifespan)

UPLOAD_DIR = Path("./uploads")
OUTPUT_DIR = Path("./output")
OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
app.mount("/output", StaticFiles(directory=str(OUTPUT_DIR)), name="output")

STATIC_DIR = Path("./static")
if STATIC_DIR.is_dir():
    app.mount("/static", StaticFiles(directory=str(STATIC_DIR)), name="static")

# All three stores mirror to disk so a restart leaves a record rather than a
# 404. Only video jobs can actually be resumed from theirs — a HeyGen render
# outlives the process — but knowing how far a batch got is worth the file.
jobs_store = JobsStore(persist_dir=OUTPUT_DIR / "batch" / "_jobs")
video_jobs_store = VideoJobsStore(persist_dir=OUTPUT_DIR / "heygen" / "_jobs")
video_batch_jobs_store = VideoBatchJobsStore(persist_dir=OUTPUT_DIR / "heygen" / "_batches")
session_config_store = SessionConfigStore()
# Pause / cancel for the video jobs and batches running now, by their id.
video_controls = JobControls()


async def _controlled(key: str, coro) -> None:
    """Run a video job or batch, dropping its control once it is over."""
    try:
        await coro
    finally:
        video_controls.close(key)

VIDEO_OUTPUT_DIR = OUTPUT_DIR / "heygen"
VIDEO_OUTPUT_DIR.mkdir(parents=True, exist_ok=True)

# Every HeyGen render, matched to its job; holds the listing and media caches.
video_library = LibraryService(
    video_jobs_store, jobs_dir=VIDEO_OUTPUT_DIR, cache_dir=VIDEO_OUTPUT_DIR / "_library",
)

# Generated scripts, kept so the writer can be shown what it already said. Not a
# job store: nothing here is in flight, and every read comes off disk.
script_history = ScriptHistoryStore(OUTPUT_DIR / "heygen" / "_scripts")
# Reads and records that history itself, both when a day is written and when
# its rows are submitted; the routes only call it.
script_writer = ScriptWriter(script_history)


async def _session_config_for_request(request: Request) -> RuntimeConfig | None:
    session_id = request.cookies.get(SESSION_COOKIE_NAME)
    if not session_id:
        return None
    return await session_config_store.get(session_id)


@dataclass(frozen=True)
class VideoDeps:
    """Everything a video job runs against, resolved at the edge."""

    renderer: VideoRenderer
    speech: SpeechSynth
    slots: TalkingPhotoSlots
    heygen: HeyGenSettings
    nas: NasConfig
    email: EmailSettings


def _build_renderer(session: RuntimeConfig | None) -> VideoRenderer:
    """Just the renderer, for requests that need HeyGen and nothing else — so a
    missing ElevenLabs or NAS setting can't break them."""
    return HeyGenRenderer(HeyGenSettings.resolve(session).api_key)


def _build_video_deps(session: RuntimeConfig | None) -> VideoDeps:
    """Resolve credentials and settings once, here, and build the seams a video
    job runs against.

    A missing key raises here — so it fails the request that asked for the work,
    instead of failing a background job several minutes in.
    """
    heygen = HeyGenSettings.resolve(session)
    eleven = ElevenLabsSettings.resolve(session)
    renderer = HeyGenRenderer(heygen.api_key)
    return VideoDeps(
        renderer=renderer,
        speech=ElevenLabsSpeech(eleven.api_key),
        slots=TalkingPhotoSlots(renderer),
        heygen=heygen,
        nas=NasConfig.resolve(session),
        email=EmailSettings.resolve(session),
    )


def _set_session_cookie(response: Response, session_id: str) -> None:
    response.set_cookie(
        key=SESSION_COOKIE_NAME,
        value=session_id,
        httponly=True,
        samesite="lax",
    )


# The HTML must always be revalidated. Without an explicit Cache-Control browsers
# apply heuristic freshness and can serve a stale page against freshly deployed
# JS — the markup then lacks the ids the new script expects and the page dies on
# load. The ?v= query on the asset tags only helps if the HTML carrying it is
# current, so the pages say no-cache and /static assets stay cacheable.
NO_HTML_CACHE = {"Cache-Control": "no-cache, must-revalidate"}


@app.get("/")
def index() -> FileResponse:
    index_path = Path("./index.html")
    if not index_path.exists():
        raise HTTPException(status_code=404, detail="index.html not found")
    return FileResponse(index_path, headers=NO_HTML_CACHE)


@app.get("/videogen")
def videogen_page() -> FileResponse:
    """Both pipelines live in one app now. This stays a real route so the Video
    section can be linked and bookmarked directly — the shell reads the path and
    opens it."""
    return index()


@app.get("/story")
def story_page() -> FileResponse:
    """The Story section's own linkable path, like /videogen."""
    return index()


@app.get("/library")
def library_page() -> FileResponse:
    """The video library's own linkable path, like /videogen."""
    return index()


@app.get("/heygen", include_in_schema=False)
def heygen_page_legacy() -> RedirectResponse:
    """The old name. Kept as a redirect rather than deleted: the video editor has
    this bookmarked, and a 308 renames the URL without costing them anything."""
    return RedirectResponse("/videogen", status_code=308)


@app.get("/video/heygen/talking-photos")
async def list_heygen_talking_photos(request: Request):
    session = await _session_config_for_request(request)
    try:
        photos = await _build_renderer(session).list_photos()
        return {"items": [
            {"talking_photo_id": p.talking_photo_id, "name": p.name, "preview_url": p.preview_url}
            for p in photos
        ]}
    except MissingSettingError as exc:
        raise HTTPException(status_code=500, detail=str(exc)) from exc
    except Exception as exc:
        raise HTTPException(status_code=502, detail=f"HeyGen list failed: {exc}") from exc


NO_PHOTO_DETAIL = "provide either an image file or a talking_photo_id"


@app.post("/video/heygen", status_code=202)
async def create_heygen_video_job(
    request: Request,
    image: UploadFile | None = File(default=None),
    talking_photo_id: str | None = Form(default=None),
    script: str = Form(...),
    character: str = Form(default="indian"),
    voice_id: str | None = Form(default=None),
    motion_prompt: str | None = Form(default=None),
    width: int | None = Form(default=None),
    height: int | None = Form(default=None),
    # video_title becomes the NAS filename (DD-MM-YYYY/<title>.mp4), so a shared
    # default silently overwrites the previous render of the day. Callers should
    # always send a distinct title; the default stays only for compatibility.
    video_title: str = Form(default="HeyGen Avatar IV Job"),
    publish_date: str | None = Form(default=None),
    stability: float = Form(default=0.5),
    similarity_boost: float = Form(default=0.75),
    style: float = Form(default=0.0),
    # Toggles arrive as the raw strings the form sent; intake parses them.
    use_speaker_boost: str | None = Form(default=None),
    # Opt-outs for the branded overlay cards and the music bed (independent).
    # Batch takes the same two; only recovery has no knobs, and it replays the
    # job's saved spec.
    include_overlay: str | None = Form(default=None),
    include_music: str | None = Form(default=None),
):
    session = await _session_config_for_request(request)

    try:
        spec = spec_from_single(SingleSubmission(
            script=script,
            video_title=video_title,
            character=character,
            voice_id=voice_id,
            motion_prompt=motion_prompt,
            publish_date=publish_date,
            talking_photo_id=talking_photo_id,
            width=width,
            height=height,
            stability=stability,
            similarity_boost=similarity_boost,
            style=style,
            use_speaker_boost=use_speaker_boost,
            include_overlay=include_overlay,
            include_music=include_music,
        ))
    except IntakeError as exc:
        raise HTTPException(status_code=400, detail=str(exc)) from exc

    # Judged on the spec, so a whitespace-only id counts as none.
    if not (spec.talking_photo_id or (image and image.filename)):
        raise HTTPException(status_code=400, detail=NO_PHOTO_DETAIL)

    image_bytes = b""
    image_filename = "image.jpg"
    if image and image.filename:
        image_bytes = await image.read()
        if not image_bytes:
            raise HTTPException(status_code=400, detail="image upload was empty")
        image_filename = image.filename

    try:
        deps = _build_video_deps(session)
    except MissingSettingError as exc:
        raise HTTPException(status_code=400, detail=str(exc)) from exc

    job_id = uuid.uuid4().hex
    await video_jobs_store.create(job_id)

    control = video_controls.open(job_id)
    spawn(
        _controlled(job_id, run_video_job(
            job_id=job_id,
            spec=spec,
            image_bytes=image_bytes,
            image_filename=image_filename,
            output_dir=VIDEO_OUTPUT_DIR,
            jobs_store=video_jobs_store,
            renderer=deps.renderer,
            speech=deps.speech,
            slots=deps.slots,
            nas_config=deps.nas,
            voice_id=deps.heygen.voice_for_character(spec.character),
            control=control,
        )),
        name=f"video-job:{job_id}",
    )

    return {"job_id": job_id, "status": "queued"}


# --- library ---------------------------------------------------------------
# Declared before /video/heygen/{job_id}, which would otherwise take "library"
# for a job id. The rules live in services/video_pipeline/library.py and the
# caching in library_service.py; these only translate to and from HTTP.

# A render's thumbnail never changes, so the browser may keep it for good.
IMMUTABLE = {"Cache-Control": "public, max-age=31536000, immutable"}


def _csv(raw: str | None) -> set[str]:
    return set((raw or "").split(",")) - {""}


def _renderer_for(session: RuntimeConfig | None) -> VideoRenderer:
    try:
        return _build_renderer(session)
    except MissingSettingError as exc:
        raise HTTPException(status_code=400, detail=str(exc)) from exc


async def _record_recovery(state: VideoJobState) -> None:
    """A recovered job's batch row, if it has one, is done now too."""
    video_path = state.summary.video_path
    await video_batch_jobs_store.row_recovered(
        state.job_id,
        nas_path=state.summary.nas_path,
        video_local_url=to_output_url(video_path, OUTPUT_DIR) if video_path else None,
    )


@app.get("/video/heygen/library")
async def get_video_library(
    request: Request,
    refresh: bool = False,
    state: str | None = None,
    q: str | None = None,
    ids: str | None = None,
    offset: int = 0,
    limit: int = 50,
):
    """Every render on the HeyGen account, newest first, each with what can be
    done with it. ``counts`` covers the whole library; ``items`` the filtered
    page. ``ids`` lets the page ask after the renders it sent, wherever they are."""
    renderer = _renderer_for(await _session_config_for_request(request))
    lib = await video_library.library(renderer, refresh=refresh)
    page = page_of(lib.entries, states=_csv(state), ids=_csv(ids), query=q or "",
                   offset=offset, limit=min(limit, 200))
    return {
        "items": [e.to_dict() for e in page.items],
        "total": page.total,
        "counts": page.counts,
        "listing_error": lib.listing_error,
    }


@app.post("/video/heygen/library/{video_id}/send-to-nas", status_code=202)
async def send_render_to_nas(video_id: str, request: Request, details: FileRenderRequest | None = None):
    """File a finished render to the NAS without re-rendering it. A render one
    of our jobs made is recovered from that job's own details; one we know
    nothing about needs ``details`` (title, publish date, character)."""
    session = await _session_config_for_request(request)
    try:
        deps = _build_video_deps(session)
    except MissingSettingError as exc:
        raise HTTPException(status_code=400, detail=str(exc)) from exc

    entry = await video_library.entry(deps.renderer, video_id)
    if entry is None:
        raise HTTPException(status_code=404, detail="No such render on HeyGen or in our jobs")

    if entry.state == "sendable":
        try:
            await start_recovery(
                job_id=entry.job_id, jobs_store=video_jobs_store, output_dir=VIDEO_OUTPUT_DIR,
                renderer=deps.renderer, nas_config=deps.nas, launch=spawn,
                on_recovered=_record_recovery,
            )
        except NotRecoverable as exc:
            raise HTTPException(status_code=409, detail=exc.reason) from exc
        return {"video_id": video_id, "job_id": entry.job_id, "status": "recovering"}

    if entry.state == "needs_details":
        if details is None:
            raise HTTPException(status_code=422, detail="This render has no title or date on file; send them with it")
        title = scrub_title(details.title)
        if not title:
            raise HTTPException(status_code=422, detail="title is empty once unsafe characters are removed")
        spec = VideoJobSpec(
            script="", video_title=title, character=details.character.strip().lower() or "indian",
            publish_date=details.publish_date, include_overlay=details.include_overlay,
            include_music=details.include_music,
        )
        job_id = await start_filing_render(
            video_id=video_id, spec=spec, jobs_store=video_jobs_store, output_dir=VIDEO_OUTPUT_DIR,
            renderer=deps.renderer, nas_config=deps.nas, launch=spawn,
        )
        return {"video_id": video_id, "job_id": job_id, "status": "recovering"}

    refusal = {
        "on_nas": f"Already on the NAS at {entry.nas_path}",
        "in_progress": "Already being worked on",
        "rendering": "Still rendering on HeyGen; try again when it finishes",
    }.get(entry.state, entry.reason)
    raise HTTPException(status_code=409, detail=refusal)


async def _download(video_id: str, request: Request) -> Download:
    renderer = _renderer_for(await _session_config_for_request(request))
    try:
        return await video_library.download(renderer, video_id)
    except NoMedia as exc:
        raise HTTPException(status_code=409, detail=str(exc)) from exc
    except (LookupError, ValueError) as exc:
        raise HTTPException(status_code=404, detail=str(exc)) from exc
    except Exception as exc:
        raise HTTPException(status_code=502, detail=f"Could not fetch the render from HeyGen: {exc}") from exc


@app.post("/video/heygen/library/{video_id}/prepare-download")
async def prepare_render_download(video_id: str, request: Request):
    """Make sure the video is on this server, then say where to fetch it. Two
    steps so the page can report a failure, then let the browser stream the
    file to disk instead of holding 30 MB in memory."""
    await _download(video_id, request)
    return {"url": f"/video/heygen/library/{video_id}/download"}


@app.get("/video/heygen/library/{video_id}/download")
async def download_render(video_id: str, request: Request):
    """The video as a file: what a job of ours filed when we have it, otherwise
    the raw render, fetched from HeyGen once and kept."""
    dl = await _download(video_id, request)
    return FileResponse(dl.path, media_type="video/mp4", filename=dl.filename)


@app.get("/video/heygen/library/{video_id}/thumbnail")
async def render_thumbnail(video_id: str, request: Request):
    """The render's thumbnail: fetched from HeyGen the first time anyone asks,
    from disk after that, and from the browser's own cache after that."""
    renderer = _renderer_for(await _session_config_for_request(request))
    try:
        path = await video_library.thumbnail(renderer, video_id)
    except (NoMedia, ValueError) as exc:
        # Not cached: a render still processing will have one later.
        raise HTTPException(status_code=404, detail=str(exc), headers={"Cache-Control": "no-store"}) from exc
    except Exception as exc:
        raise HTTPException(status_code=502, detail=f"Could not fetch the thumbnail from HeyGen: {exc}",
                            headers={"Cache-Control": "no-store"}) from exc
    return FileResponse(path, media_type=sniff_image_type(path), headers=IMMUTABLE)


@app.get("/video/heygen/{job_id}")
async def get_heygen_video_job(job_id: str):
    state = await video_jobs_store.get(job_id)
    if not state:
        raise HTTPException(status_code=404, detail="Job not found")
    payload = state.model_dump(mode="json")
    if state.summary.video_path:
        rel = to_output_url(state.summary.video_path, OUTPUT_DIR)
        payload["video_local_url"] = rel
    control = video_controls.get(job_id)
    payload["cancel_requested"] = bool(control and control.cancelled)
    return payload


@app.post("/video/heygen/{job_id}/cancel")
async def cancel_heygen_video_job(job_id: str):
    """Stop a single render at its next step. Once HeyGen has it, the render
    still finishes (and is billed) on their side, but it is not filed."""
    if not await video_jobs_store.get(job_id):
        raise HTTPException(status_code=404, detail="Job not found")
    control = video_controls.get(job_id)
    if control is None:
        raise HTTPException(status_code=409, detail="Job is not running, so there is nothing to cancel")
    control.cancel()
    return {"job_id": job_id, "status": "cancelling"}


@app.get("/video/heygen/jobs/recoverable")
async def list_recoverable_heygen_jobs():
    """Failed jobs whose HeyGen render actually finished — re-runnable via recover."""
    ids = await video_jobs_store.list_recoverable()
    return {"recoverable": ids, "count": len(ids)}


@app.post("/video/heygen/{job_id}/recover", status_code=202)
async def recover_heygen_video_job(job_id: str, request: Request):
    """Re-download a finished-but-failed render and push it to the NAS. Recovers
    a job that died on the download/NAS step (its HeyGen render is intact)."""
    session = await _session_config_for_request(request)
    try:
        deps = _build_video_deps(session)
    except MissingSettingError as exc:
        raise HTTPException(status_code=400, detail=str(exc)) from exc

    try:
        await start_recovery(
            job_id=job_id,
            jobs_store=video_jobs_store,
            output_dir=VIDEO_OUTPUT_DIR,
            renderer=deps.renderer,
            nas_config=deps.nas,
            launch=spawn,
            on_recovered=_record_recovery,
        )
    except NotRecoverable as exc:
        status = 404 if exc.kind == "not_found" else 409
        raise HTTPException(status_code=status, detail=exc.reason) from exc
    return {"job_id": job_id, "status": "recovering"}


@app.post("/video/heygen/recover-failed", status_code=202)
async def recover_all_failed_heygen_jobs(request: Request):
    """Recover every failed job whose render finished on HeyGen, in one shot.
    Jobs already being recovered are passed by."""
    session = await _session_config_for_request(request)
    try:
        deps = _build_video_deps(session)
    except MissingSettingError as exc:
        raise HTTPException(status_code=400, detail=str(exc)) from exc

    ids = await start_recovery_of_failed(
        jobs_store=video_jobs_store,
        output_dir=VIDEO_OUTPUT_DIR,
        renderer=deps.renderer,
        nas_config=deps.nas,
        launch=spawn,
        on_recovered=_record_recovery,
    )
    started = set(ids)
    video_ids = [j.summary.heygen_video_id for j in await video_jobs_store.snapshot() if j.job_id in started]
    return {"status": "recovering", "job_ids": ids, "video_ids": video_ids, "count": len(ids)}


@app.post("/video/scripts/generate")
async def generate_video_scripts(
    request: Request,
    brief: str = Form(...),
    language: str = Form(default="hi-IN"),
    publish_date: str = Form(...),
    category: str = Form(default="default"),
    item_set: str = Form(default="zodiac"),
    title: str | None = Form(default=None),
):
    """Write the day's scripts and hand them back as drafts.

    Nothing is rendered here and no job is created: the response is text for an
    operator to read and edit. A render costs money at submission, so the two
    steps stay separate — see services/script_writer.py.

    Synchronous rather than a background job. One Gemini call for the whole set
    takes seconds, and a draft the operator is waiting to read has nothing to
    gain from a polling endpoint.
    """
    session = await _session_config_for_request(request)

    try:
        settings = ScriptWriterSettings.resolve(session)
    except MissingSettingError as exc:
        raise HTTPException(status_code=400, detail=str(exc)) from exc

    # Blocking throughout — history read, the Gemini call, history write — so
    # the whole day runs off the event loop. A set whose repair budget ran out
    # still comes back, with the unresolved clashes as warnings.
    try:
        day = await asyncio.to_thread(
            script_writer.write_day,
            category=category,
            language=language,
            publish_date=publish_date,
            brief=brief,
            item_set=item_set,
            title=title,
            settings=settings,
        )
    except ScriptRequestError as exc:
        raise HTTPException(status_code=400, detail=str(exc)) from exc
    except ScriptWriterError as exc:
        raise HTTPException(status_code=502, detail=str(exc)) from exc

    return {
        "language": language,
        "publish_date": publish_date,
        "items": [{"video_title": d.title, "script": d.script} for d in day.drafts],
        "warnings": day.warnings,
    }


@app.post("/video/heygen/batch", status_code=202)
async def create_heygen_batch_job(
    request: Request,
    # Either an uploaded image or a photo picked from the account, as for a
    # single video.
    image: UploadFile | None = File(default=None),
    talking_photo_id: str | None = Form(default=None),
    excel: UploadFile | None = File(default=None),
    rows: str | None = Form(default=None),
    character: str = Form(default="indian"),
    motion_prompt: str | None = Form(default=None),
    publish_date: str | None = Form(default=None),
    # Only sent when the rows were written by the script writer. They identify
    # which history these rows belong to, so what is submitted can replace what
    # was generated — see the recording step below.
    category: str | None = Form(default=None),
    language: str | None = Form(default=None),
    # Strings, not bools: intake parses them, as for a single video.
    include_overlay: str | None = Form(default=None),
    include_music: str | None = Form(default=None),
):
    session = await _session_config_for_request(request)

    # Two ways in, one batch: a spreadsheet the operator authored, or rows of
    # generated scripts they reviewed in the browser. Both become the same row
    # list before anything else happens, so nothing downstream knows which.
    has_excel = bool(excel and excel.filename)
    has_rows = bool(rows and rows.strip())
    if has_excel == has_rows:
        raise HTTPException(
            status_code=400, detail="provide either an excel file or rows JSON, not both"
        )
    if has_excel:
        ensure_file_extension(excel.filename, ".xlsx", "Only .xlsx files are allowed")
        excel_tmp = tempfile.NamedTemporaryFile(delete=False, suffix=".xlsx")
        excel_path = Path(excel_tmp.name)
        excel_tmp.close()
        await save_upload_file(excel, excel_path)
        try:
            batch_rows_parsed = await asyncio.to_thread(read_heygen_batch_rows, excel_path)
        except BatchExcelError as exc:
            raise HTTPException(status_code=400, detail=str(exc)) from exc
        finally:
            excel_path.unlink(missing_ok=True)
        empty_detail = "Excel has no non-empty script rows"
    else:
        try:
            batch_rows_parsed = read_heygen_batch_rows_json(rows or "")
        except BatchExcelError as exc:
            raise HTTPException(status_code=400, detail=str(exc)) from exc
        empty_detail = "No non-empty scripts were submitted"

    if not batch_rows_parsed:
        raise HTTPException(status_code=400, detail=empty_detail)

    # Before history is recorded, so a batch refused here leaves no trace there.
    try:
        items = specs_from_batch(batch_rows_parsed, BatchSettings(
            character=character,
            motion_prompt=motion_prompt,
            publish_date=publish_date,
            talking_photo_id=talking_photo_id,
            include_overlay=include_overlay,
            include_music=include_music,
        ))
    except IntakeError as exc:
        raise HTTPException(status_code=400, detail=str(exc)) from exc

    # Intake cleaned the id (blank to None) and put it on every row alike.
    picked_photo_id = items[0].spec.talking_photo_id
    if not (picked_photo_id or (image and image.filename)):
        raise HTTPException(status_code=400, detail=NO_PHOTO_DETAIL)

    image_bytes = b""
    image_filename = "image.jpg"
    if not picked_photo_id:
        image_bytes = await image.read()
        if not image_bytes:
            raise HTTPException(status_code=400, detail="image upload was empty")
        image_filename = image.filename

    # Record what is actually being rendered, replacing the drafts recorded when
    # they were written — the operator's edits are what tomorrow must avoid.
    # Only for generated rows: a spreadsheet carries no category or language to
    # file it under, and its scripts were not written by the model that would be
    # told to avoid them.
    if has_rows:
        try:
            await asyncio.to_thread(
                script_writer.record_submitted,
                category=category,
                language=language,
                publish_date=publish_date,
                rows=[(item.spec.video_title, item.spec.script) for item in items],
            )
        except UnknownItemTitleError as exc:
            raise HTTPException(status_code=400, detail=str(exc)) from exc

    try:
        deps = _build_video_deps(session)
    except MissingSettingError as exc:
        raise HTTPException(status_code=400, detail=str(exc)) from exc

    batch_id = uuid.uuid4().hex
    batch_rows = [
        BatchRowState(row_index=item.row_index, script=item.spec.script, video_title=item.spec.video_title)
        for item in items
    ]
    await video_batch_jobs_store.create(batch_id, batch_rows)

    control = video_controls.open(batch_id)
    spawn(
        _controlled(batch_id, run_video_batch_job(
            batch_id=batch_id,
            items=items,
            image_bytes=image_bytes,
            image_filename=image_filename,
            talking_photo_id=picked_photo_id,
            output_dir=VIDEO_OUTPUT_DIR,
            output_base_dir=OUTPUT_DIR,
            batch_store=video_batch_jobs_store,
            video_jobs_store=video_jobs_store,
            renderer=deps.renderer,
            speech=deps.speech,
            slots=deps.slots,
            heygen=deps.heygen,
            nas_config=deps.nas,
            email=deps.email,
            control=control,
        )),
        name=f"video-batch:{batch_id}",
    )

    return {"batch_id": batch_id, "status": "queued", "total": len(items)}


@app.get("/video/heygen/batch/{batch_id}")
async def get_heygen_batch_job(batch_id: str):
    state = await video_batch_jobs_store.get(batch_id)
    if not state:
        raise HTTPException(status_code=404, detail="Batch job not found")
    control = video_controls.get(batch_id)
    if control is not None:
        state.paused = control.paused
        state.cancel_requested = control.cancelled
    return state.model_dump(mode="json")


async def _batch_control(batch_id: str):
    if not await video_batch_jobs_store.get(batch_id):
        raise HTTPException(status_code=404, detail="Batch job not found")
    control = video_controls.get(batch_id)
    if control is None:
        raise HTTPException(status_code=409, detail="Batch is not running")
    return control


@app.post("/video/heygen/batch/{batch_id}/pause")
async def pause_heygen_batch(batch_id: str):
    """Start no new row, and hold rows HeyGen does not have yet. Renders already
    submitted keep going: they are paid for."""
    control = await _batch_control(batch_id)
    if control.cancelled:
        raise HTTPException(status_code=409, detail="Batch is already being cancelled")
    control.pause()
    return {"batch_id": batch_id, "paused": True}


@app.post("/video/heygen/batch/{batch_id}/resume")
async def resume_heygen_batch(batch_id: str):
    control = await _batch_control(batch_id)
    control.resume()
    return {"batch_id": batch_id, "paused": False}


@app.post("/video/heygen/batch/{batch_id}/cancel")
async def cancel_heygen_batch(batch_id: str):
    """Stop every row that has not finished. Rows already on the NAS stay."""
    control = await _batch_control(batch_id)
    control.cancel()
    return {"batch_id": batch_id, "status": "cancelling"}


@app.get("/health")
async def health_check() -> dict:
    """Health check endpoint for Docker healthchecks and monitoring."""
    return {"status": "healthy", "service": "autodub"}


@app.post("/config/session-env", response_model=SessionEnvConfigResponse)
async def set_session_env_config(payload: SessionEnvConfigRequest, request: Request):
    parsed = parse_env_text(payload.env_text)
    # Required keys are derived from the settings classes, so this check can no
    # longer drift from what the app actually reads.
    absent = [key for key in required_keys() if not parsed.get(key, "").strip()]
    if absent:
        raise HTTPException(
            status_code=400,
            detail={
                "message": "Missing required keys in pasted .env text",
                "missing_keys": absent,
            },
        )

    session_id = request.cookies.get(SESSION_COOKIE_NAME) or session_config_store.generate_session_id()
    await session_config_store.set(session_id, parsed)

    response = JSONResponse(
        status_code=200,
        content=SessionEnvConfigResponse(
            configured=True, missing_keys=[], required_keys=required_keys()
        ).model_dump(),
    )
    _set_session_cookie(response, session_id)
    return response


@app.get("/config/session-env/status", response_model=SessionEnvConfigResponse)
async def get_session_env_config_status(request: Request):
    session = await _session_config_for_request(request)
    absent = missing_keys(session)
    return SessionEnvConfigResponse(
        configured=not absent, missing_keys=absent, required_keys=required_keys()
    )


@app.delete("/config/session-env", response_model=SessionEnvConfigResponse)
async def clear_session_env_config(request: Request):
    session_id = request.cookies.get(SESSION_COOKIE_NAME)
    if session_id:
        await session_config_store.clear(session_id)
    response = JSONResponse(
        status_code=200,
        content=SessionEnvConfigResponse(
            configured=False, missing_keys=required_keys(), required_keys=required_keys()
        ).model_dump(),
    )
    response.delete_cookie(SESSION_COOKIE_NAME)
    return response


@app.get("/logs/important")
async def get_important_logs(since_id: int = 0, limit: int = 200) -> dict:
    return _get_important_logs(since_id, limit)


@app.post("/translate")
async def translate_pipeline(payload: TranslateRequest, request: Request):
    session = await _session_config_for_request(request)
    try:
        sarvam = SarvamSettings.resolve(session)
    except MissingSettingError as exc:
        raise HTTPException(status_code=400, detail=str(exc)) from exc
    google_translate = GoogleTranslateSettings.resolve(session)

    transcript = payload.text
    translations = {}
    for lang in payload.target_languages:
        try:
            translated = translate_with_fallback(
                transcript,
                settings=sarvam,
                google_translate=google_translate,
                target_language_code=lang,
                source_language_code="auto",
            )
        except Exception as exc:
            translated = f"[Translation error: {exc}]"
        translations[lang] = translated

    return {
        "input_text": transcript,
        "translations": translations,
    }


@app.post("/batch/excel-jobs", status_code=202, response_model=CreateJobResponse)
async def create_excel_job(
    request: Request,
    file: UploadFile = File(...),
    target_languages: list[str] | None = Form(default=None),
    target_languages_json: str | None = Form(default=None),
    max_language_parallelism: int | None = Form(default=None),
    teaching_mode: bool | None = Form(default=False),
    mode: str = Form(default="create"),
):
    session = await _session_config_for_request(request)
    filename = ensure_file_extension(file.filename, ".xlsx", "Only .xlsx files are allowed")
    if max_language_parallelism is not None and max_language_parallelism < 1:
        raise HTTPException(status_code=400, detail="max_language_parallelism must be >= 1")

    if mode not in {"create", "append"}:
        raise HTTPException(status_code=400, detail="mode must be 'create' or 'append'")

    parsed_languages = parse_target_languages(target_languages, target_languages_json)

    # Resolve everything the batch will need before accepting it. An Excel batch
    # runs for hours; a missing key should cost one HTTP response, not an hour.
    try:
        settings = Settings.resolve(session)
        english_voice_is_required(parsed_languages, settings.eleven)
    except (MissingSettingError, ValueError) as exc:
        raise HTTPException(status_code=400, detail=str(exc)) from exc

    temp_file = tempfile.NamedTemporaryFile(delete=False, suffix=".xlsx")
    saved_path = Path(temp_file.name)
    temp_file.close()

    await save_upload_file(file, saved_path)

    job_id = uuid.uuid4().hex
    await jobs_store.create(job_id)

    spawn(
        run_excel_batch_job(
            job_id=job_id,
            excel_path=str(saved_path),
            target_languages=parsed_languages,
            max_language_parallelism=max_language_parallelism,
            jobs_store=jobs_store,
            settings=settings,
            teaching_mode=teaching_mode,
            output_dir=OUTPUT_DIR,
            mode=mode,
        ),
        name=f"excel-batch:{job_id}",
    )

    return CreateJobResponse(job_id=job_id, status="queued")


@app.post("/batch/story/split")
async def split_story_route(payload: StorySplitRequest, request: Request):
    """Break a pasted story into voiceover lines with emotions, for review.
    Synchronous: one Gemini call (two if the first split fails validation)."""
    if not payload.text.strip():
        raise HTTPException(status_code=400, detail="Paste a story first: the text is empty")
    session = await _session_config_for_request(request)
    try:
        settings = QCSettings.resolve(session)
    except (MissingSettingError, ValueError) as exc:
        raise HTTPException(status_code=400, detail=str(exc)) from exc

    try:
        lines = await asyncio.to_thread(split_story, payload.text, settings=settings)
    except StorySplitError as exc:
        raise HTTPException(status_code=400, detail=str(exc)) from exc
    except StoryGeminiError as exc:
        raise HTTPException(status_code=502, detail=str(exc)) from exc

    return {
        "activity_name": story_activity_name(payload.title),
        "lines": [{"text": line.text, "emotion": line.emotion} for line in lines],
        # The review table is fed from here so it and the validator share one
        # list and one cap.
        "emotions": STORY_EMOTIONS,
        "max_line_chars": MAX_LINE_CHARS,
    }


@app.post("/batch/story/xlsx")
async def story_xlsx_route(payload: StoryXlsxRequest):
    """The reviewed lines as a sheet the normal Audio (Excel) flow accepts."""
    from openpyxl import Workbook

    activity = payload.activity_name.strip()
    if not activity:
        raise HTTPException(status_code=400, detail="The activity name is empty")
    if not payload.lines:
        raise HTTPException(status_code=400, detail="There are no lines to download")
    for n, line in enumerate(payload.lines, start=1):
        if not line.text.strip():
            raise HTTPException(status_code=400, detail=f"Line {n} is empty")
        if line.emotion not in STORY_EMOTIONS:
            raise HTTPException(status_code=400, detail=f"Line {n} has an unknown emotion: {line.emotion}")

    workbook = Workbook()
    sheet = workbook.active
    sheet.append(REQUIRED_HEADERS)
    for n, line in enumerate(payload.lines, start=1):
        sheet.append([line.text.strip(), line.emotion, activity, f"{activity}{n}"])
    buffer = BytesIO()
    workbook.save(buffer)
    filename = re.sub(r"[^A-Za-z0-9_.-]", "_", activity) + ".xlsx"
    return Response(
        content=buffer.getvalue(),
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f'attachment; filename="{filename}"'},
    )


@app.post("/batch/story-jobs", status_code=202, response_model=CreateJobResponse)
async def create_story_job(payload: StoryJobRequest, request: Request):
    """Voice a reviewed story: translate, one story fix per language, speak.
    Polled and cancelled through the Excel job endpoints."""
    activity = payload.activity_name.strip()
    if not activity:
        raise HTTPException(status_code=400, detail="The activity name is empty")
    if not payload.lines:
        raise HTTPException(status_code=400, detail="There are no lines to voice")
    lines = [StoryLine(text=line.text.strip(), emotion=line.emotion) for line in payload.lines]
    for n, line in enumerate(lines, start=1):
        problem = line_problem(line)
        if problem:
            raise HTTPException(status_code=400, detail=f"Line {n}: {problem}")

    parsed_languages = parse_target_languages(payload.target_languages, None)

    session = await _session_config_for_request(request)
    try:
        settings = Settings.resolve(session)
        english_voice_is_required(parsed_languages, settings.eleven)
    except (MissingSettingError, ValueError) as exc:
        raise HTTPException(status_code=400, detail=str(exc)) from exc

    job_id = uuid.uuid4().hex
    await jobs_store.create(job_id)

    spawn(
        run_story_job(
            job_id=job_id,
            title=payload.title.strip(),
            activity_name=activity,
            lines=lines,
            target_languages=parsed_languages,
            jobs_store=jobs_store,
            settings=settings,
            output_dir=OUTPUT_DIR,
        ),
        name=f"story-job:{job_id}",
    )

    return CreateJobResponse(job_id=job_id, status="queued")


@app.get("/batch/excel-jobs/{job_id}", response_model=JobState)
async def get_excel_job(job_id: str):
    job = await jobs_store.get(job_id)
    if not job:
        raise HTTPException(status_code=404, detail="Job not found")
    return job


@app.post("/batch/excel-jobs/{job_id}/cancel")
async def cancel_excel_job(job_id: str):
    job = await jobs_store.get(job_id)
    if not job:
        raise HTTPException(status_code=404, detail="Job not found")
    ok = await jobs_store.request_cancel(job_id)
    if not ok:
        raise HTTPException(status_code=409, detail=f"Job is {job.status}, cannot cancel")
    return {"job_id": job_id, "status": "cancelling"}


@app.post("/batch/preview-excel")
async def preview_excel(file: UploadFile = File(...)):
    ensure_file_extension(file.filename, ".xlsx", "Only .xlsx files are allowed")
    import tempfile, openpyxl
    tmp = tempfile.NamedTemporaryFile(delete=False, suffix=".xlsx")
    saved = Path(tmp.name)
    tmp.close()
    await save_upload_file(file, saved)
    try:
        wb = openpyxl.load_workbook(str(saved), read_only=True, data_only=True)
        ws = wb.worksheets[0]
        rows_out = []
        for i, row in enumerate(ws.iter_rows(values_only=True)):
            cells = [str(c).strip() if c is not None else "" for c in row]
            rows_out.append(cells)
            if i >= 3:
                break
        wb.close()
        return {"rows": rows_out}
    finally:
        saved.unlink(missing_ok=True)


@app.post("/stt")
async def stt_pipeline(
    request: Request,
    audio: UploadFile = File(...),
    target_language: str = Form("hi-IN"),
) -> dict:
    session = await _session_config_for_request(request)
    try:
        sarvam = SarvamSettings.resolve(session)
    except MissingSettingError as exc:
        raise HTTPException(status_code=400, detail=str(exc)) from exc
    google_translate = GoogleTranslateSettings.resolve(session)
    filename = ensure_file_extension(audio.filename, ".mp3", "Only .mp3 files are allowed")

    UPLOAD_DIR.mkdir(parents=True, exist_ok=True)
    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)

    safe_stem_ = safe_stem(filename)
    upload_path = UPLOAD_DIR / f"{safe_stem_}-{uuid.uuid4().hex}.mp3"

    await save_upload_file(audio, upload_path)

    transcripts = transcribe_audio(
        audio_paths=[str(upload_path)],
        settings=sarvam,
        output_dir=str(OUTPUT_DIR),
        language_code="unknown",
    )

    transcript = next(iter(transcripts.values()))
    translated = translate_with_fallback(
        transcript,
        settings=sarvam,
        google_translate=google_translate,
        target_language_code=target_language,
        source_language_code="auto",
    )

    output_path = OUTPUT_DIR / f"{upload_path.stem}.{target_language}"
    tts_result = text_to_speech(
        translated,
        target_language_code=target_language,
        output_path=str(output_path),
        settings=sarvam,
        speaker="shubh",
    )

    tts_url = to_output_url(tts_result, OUTPUT_DIR)

    return {
        "input_file": filename,
        "transcript": transcript,
        "translation": translated,
        "tts_output": tts_result,
        "tts_url": tts_url,
    }


@app.post("/finalize-text")
async def finalize_text(request: FinalizeTextRequest):
    logging.info("Finalize text for language %s: %s", request.language, request.text)
    return {"status": "ok", "language": request.language, "text": request.text}


@app.post("/tts-elevenlabs")
async def tts_elevenlabs(payload: ElevenLabsTTSRequest, request: Request):
    session = await _session_config_for_request(request)
    try:
        eleven = ElevenLabsSettings.resolve(session)
    except MissingSettingError as exc:
        raise HTTPException(status_code=400, detail=str(exc)) from exc
    try:
        audio_bytes = synthesize_speech_bytes(
            payload.text,
            api_key=eleven.api_key,
            config=ElevenLabsTTSConfig(
                voice_id=payload.voice_id,
                model_id=payload.model_id,
                stability=payload.stability,
                similarity_boost=payload.similarity_boost,
                style=payload.style,
                use_speaker_boost=payload.use_speaker_boost,
            ),
        )
    except Exception as exc:
        raise HTTPException(status_code=500, detail=f"ElevenLabs TTS failed: {exc}") from exc

    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    out_path = OUTPUT_DIR / f"elevenlabs-{uuid.uuid4().hex}.mp3"
    try:
        out_path.write_bytes(audio_bytes)
    except OSError as exc:
        raise HTTPException(status_code=500, detail="Failed to save generated audio") from exc

    tts_url = to_output_url(out_path, OUTPUT_DIR)
    if not tts_url:
        raise HTTPException(status_code=500, detail="Failed to generate output URL")

    return {
        "tts_url": tts_url,
        "output_file": str(out_path),
    }
